"""Стиль автора по образцам (#134): collecting Образцы, the result and its buttons, on fakes."""

from __future__ import annotations

import time
import zipfile
from io import BytesIO

import pytest
from docx import Document

from content_zavod.pipelines import ANALYZE_STYLE_JOB as WORKER_JOB
from content_zavod.settings import OwnerSettings, SettingsService, StyleProfile
from content_zavod.telegram.input_prompt import InputPrompt
from content_zavod.telegram.settings_screen import SETTING_FIELDS
from content_zavod.telegram.style_samples import (
    ANALYZE_STYLE_JOB,
    MAX_FILE_BYTES,
    STYLE_FIX_KIND,
    STYLE_SAMPLES_KIND,
    IncomingSample,
    SampleDocument,
    StyleRef,
    StyleSamples,
    document_text,
    render_result,
    sample_link,
)
from content_zavod.telegram.texts import (
    STALE_BUTTON,
    STYLE_FAILED,
    STYLE_NOTHING_YET,
    STYLE_RESTARTED,
    STYLE_WARN_DZEN,
    STYLE_WARN_FILE_BIG,
    STYLE_WARN_FILE_SIZE_UNKNOWN,
    STYLE_WARN_FILE_TYPE,
    STYLE_WARN_FILE_UNREADABLE,
    STYLE_WARN_LINKS,
    STYLE_WARN_TELEGRAM_LINK,
    STYLE_WORKING,
    error_alert_text,
)
from content_zavod.workspaces import DailyFreeJobLimitReached, GenerationBusy, Workspace
from tests.tenancy import WS_A, WS_B

from .fakes import (
    FakeJobs,
    FakeLimits,
    FakePendingInputs,
    RecordingBot,
    button_data,
    button_texts,
)

CHAT = USER = 42
PART = ("Я пишу коротко и по делу, без воды и канцелярита. " * 12).strip()  # 599 chars
PROFILE_OUTPUT = {
    "status": "ok",
    "profile": {
        "traits": {"tone": "прямой", "address": "на «вы»"},
        "summary": "Вы пишете коротко и по делу.",
        "voice_examples": ["без воды и канцелярита"],
        "confidence": "high",
    },
    "chars": 4000,
    "unread": [],
    "source": "articles",
}


class WorkspaceStore:
    def __init__(self) -> None:
        self.values: dict[tuple[int, str], str] = {}

    async def get(self, workspace_id: int, key: str) -> str | None:
        return self.values.get((workspace_id, key))

    async def set(self, workspace_id: int, key: str, value: str) -> None:
        self.values[(workspace_id, key)] = value


class Onboarding:
    def __init__(self, needed: bool = False) -> None:
        self.is_needed = needed

    async def needed(self, workspace_id: int) -> bool:
        return self.is_needed


class CountingPending(FakePendingInputs):
    def __init__(self) -> None:
        super().__init__()
        self.puts = 0

    async def put(self, workspace_id, chat_id, user_id, pending):
        self.puts += 1
        return await super().put(workspace_id, chat_id, user_id, pending)


class BrokenJobs(FakeJobs):
    async def enqueue(self, workspace_id, job_type, payload, idempotency_key, *, conn=None):
        raise RuntimeError("database is down")


class Env:
    def __init__(self, jobs: FakeJobs | None = None) -> None:
        self.bot = RecordingBot()
        self.pending = CountingPending()
        self.jobs = jobs or FakeJobs()
        self.store = WorkspaceStore()
        self.settings = SettingsService(self.store)
        self.files: dict[str, bytes] = {}
        self.downloaded: list[str] = []
        self.onboarding = Onboarding()
        self.limits = FakeLimits()
        self.style = self.new_style()
        self.a = Workspace(WS_A, USER, None, "Europe/Moscow", None, None)
        self.b = Workspace(WS_B, USER, None, "Europe/Moscow", None, None)

    def new_style(self) -> StyleSamples:
        """A fresh `StyleSamples` over the same waits and Jobs - the bot after a restart."""
        return StyleSamples(
            self.jobs,
            self.settings,
            self.pending,
            InputPrompt(self.bot, self.pending),
            self.bot,
            self.download,
            self.onboarding,
            self.limits,
        )

    async def download(self, file_id: str) -> bytes:
        self.downloaded.append(file_id)
        return self.files[file_id]

    async def send(self, text: str, workspace: Workspace | None = None) -> bool:
        return await self.style.handle_message(
            workspace or self.a, CHAT, USER, IncomingSample(text=text)
        )

    async def send_file(self, name: str, data: bytes | None = None, size: int | None = -1) -> bool:
        if data is not None:
            self.files[name] = data
        file_size = len(data or b"") if size == -1 else size
        document = SampleDocument(name, name, file_size)
        return await self.style.handle_message(
            self.a, CHAT, USER, IncomingSample(text=None, document=document)
        )

    async def done(self) -> str | None:
        return await self.style.done(self.a, CHAT, USER, self.question_id)

    @property
    def question_id(self) -> int:
        return self.pending.rows[(CHAT, USER)].prompt_message_id

    @property
    def last_edit(self) -> str:
        return self.bot.edited[-1][2]


@pytest.fixture
def env() -> Env:
    return Env()


def _docx(text: str) -> bytes:
    document = Document()
    for paragraph in text.split("\n"):
        document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_bot_and_worker_name_the_same_job() -> None:
    assert ANALYZE_STYLE_JOB == WORKER_JOB


# --- collecting until «Готово» ---


async def test_long_text_in_several_messages_is_collected_until_done(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    question = env.question_id
    (_, ask, markup, _) = env.bot.sent[-1]
    assert "Пришлите 1–3 свои статьи" in ask
    assert button_data(markup) == [["yd:a:s"], [f"ci:{STYLE_SAMPLES_KIND}"]]

    for _ in range(3):
        assert await env.send(PART)
    assert "Получено: 1797 знаков текста" in env.last_edit

    assert await env.done() is None

    (job,) = env.jobs.jobs.values()
    assert job.job_type == ANALYZE_STYLE_JOB
    assert job.payload["samples"] == [PART] * 3
    assert (job.payload["source"], job.payload["message_id"]) == ("articles", question)
    assert env.bot.edited[-1][1:3] == (question, STYLE_WORKING)
    assert (CHAT, USER) not in env.pending.rows
    assert not await env.send(PART)  # the wait is closed: no longer an Образец


async def test_every_accepted_sample_renews_the_wait(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    before = env.pending.puts

    await env.send(PART)
    await env.send(PART)

    assert env.pending.puts == before + 2
    assert env.pending.rows[(CHAT, USER)].kind == STYLE_SAMPLES_KIND


async def test_done_typed_as_a_word_is_the_button(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    for _ in range(3):
        await env.send(PART)

    assert await env.send("Готово!")

    (job,) = env.jobs.jobs.values()
    assert "Готово" not in "".join(job.payload["samples"])


async def test_under_1500_chars_no_job_and_the_presets_are_offered(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:o")
    await env.send(PART)

    assert await env.done() is None

    assert env.jobs.jobs == {}
    assert "Маловато текста: 599 знаков" in env.last_edit
    assert ["yp:o"] in button_data(env.bot.edited[-1][3])
    assert await env.send(PART * 2)  # still collecting
    await env.done()
    assert len(env.jobs.jobs) == 1


async def test_done_with_nothing_sent_says_so(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")

    assert await env.done() == STYLE_NOTHING_YET
    assert env.pending.rows[(CHAT, USER)].kind == STYLE_SAMPLES_KIND


async def test_a_stale_done_starts_nothing(env: Env) -> None:
    assert await env.style.done(env.a, CHAT, USER, 77) == STALE_BUTTON
    assert env.jobs.jobs == {}


async def test_a_failed_enqueue_keeps_the_samples_and_the_wait() -> None:
    env = Env(BrokenJobs())
    await env.style.start(env.a, CHAT, USER, "a:s")
    for _ in range(3):
        await env.send(PART)

    with pytest.raises(RuntimeError):
        await env.done()

    assert env.pending.rows[(CHAT, USER)].kind == STYLE_SAMPLES_KIND
    assert STYLE_WORKING not in [text for _, _, text, _ in env.bot.edited]
    env.jobs = FakeJobs()
    env.style._jobs = env.jobs
    await env.done()
    assert next(iter(env.jobs.jobs.values())).payload["samples"] == [PART] * 3


async def test_after_a_restart_the_owner_is_asked_to_send_again(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    for _ in range(3):
        await env.send(PART)
    env.style = env.new_style()  # the wait survived, the texts did not

    assert await env.done() == STYLE_RESTARTED
    assert env.jobs.jobs == {}

    env.style = env.new_style()
    await env.send(PART)
    assert STYLE_RESTARTED in env.last_edit


async def test_collections_are_forgotten_with_the_waits_ttl(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    await env.send(PART)
    env.style._collected[(WS_A, CHAT, USER)].touched = time.monotonic() - 3600

    await env.style.start(env.b, 7, 7, "a:s")

    assert (WS_A, CHAT, USER) not in env.style._collected


async def test_posts_ask_for_consent_and_are_sent_as_posts(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "p:s")
    (_, ask, markup, _) = env.bot.sent[-1]
    assert "вы подтверждаете, что канал ваш" in ask
    assert button_texts(markup)[0] == ["✅ Это мой канал — готово"]

    await env.send("t.me/my_channel")
    await env.done()

    (job,) = env.jobs.jobs.values()
    assert job.payload["links"] == ["https://t.me/my_channel"] and job.payload["samples"] == []
    assert (job.payload["source"], job.payload["mode"]) == ("posts", "p:s")


async def test_commands_are_not_samples(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")

    assert not await env.send("/menu")


def test_a_bare_mode_from_an_old_button_comes_from_settings() -> None:
    assert StyleRef.decode("p") == StyleRef("p", "s")
    assert StyleRef.decode("x:o") == StyleRef("a", "o")


# --- links ---


@pytest.mark.parametrize(
    ("text", "link"),
    [
        ("t.me/my_channel", "https://t.me/my_channel"),
        ("https://t.me/s/my_channel", "https://t.me/my_channel"),
        ("@my_channel", "https://t.me/my_channel"),
        ("vc.ru/u/1-statya", "https://vc.ru/u/1-statya"),
        ("Как я считаю рекламу\nhttps://vc.ru/u/1-statya", "https://vc.ru/u/1-statya"),
        ("Привет", None),
        ("одна ссылка vc.ru/x и текст", None),
        ("Две: https://a.ru/1 и https://b.ru/2", None),
        ("Длинный текст " * 30 + "https://vc.ru/u/1", None),
    ],
)
def test_sample_link(text: str, link: str | None) -> None:
    assert sample_link(text) == link


async def test_a_dzen_link_is_refused_at_once(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")

    await env.send("https://dzen.ru/a/aBcD")

    assert STYLE_WARN_DZEN in env.last_edit
    await env.send(PART * 3)
    await env.done()
    assert next(iter(env.jobs.jobs.values())).payload["links"] == []


async def test_a_telegram_link_for_articles_points_to_posts(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")

    await env.send("t.me/my_channel")

    assert STYLE_WARN_TELEGRAM_LINK in env.last_edit
    assert "ссылок" not in env.last_edit


async def test_at_most_three_links(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    for index in range(4):
        await env.send(f"https://blog.example/{index}")

    assert STYLE_WARN_LINKS in env.last_edit
    await env.done()
    assert len(next(iter(env.jobs.jobs.values())).payload["links"]) == 3


# --- files ---


async def test_a_docx_file_is_read(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")

    await env.send_file("статья.docx", _docx("Первый абзац.\nВторой абзац."))

    assert "Получено: 27 знаков" in env.last_edit


def test_an_inflated_docx_is_refused_without_parsing() -> None:
    """91 KB that unpacks to ~19 MB once stalled the bot for a minute."""
    paragraph = "<w:p><w:r><w:t>" + "а" * 1000 + "</w:t></w:r></w:p>"
    xml = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body>" + paragraph * 10_000 + "</w:body></w:document>"
    )
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", xml)
    data = buffer.getvalue()
    assert len(data) < MAX_FILE_BYTES

    started = time.monotonic()
    assert document_text(data, "bomb.docx") is None
    assert time.monotonic() - started < 0.5


def test_document_text_reads_txt_md_in_utf8_utf16_or_cp1251() -> None:
    assert document_text("Привет".encode(), "a.txt") == "Привет"
    assert document_text("Привет".encode("utf-16"), "a.txt") == "Привет"
    assert document_text("Привет".encode("cp1251"), "a.md") == "Привет"
    assert document_text(b"not a zip", "a.docx") is None


async def test_unsupported_too_big_unsized_or_broken_files_are_refused(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")

    await env.send_file("статья.pdf", b"%PDF")
    assert STYLE_WARN_FILE_TYPE in env.last_edit
    await env.send_file("big.txt", size=MAX_FILE_BYTES + 1)
    assert STYLE_WARN_FILE_BIG in env.last_edit
    await env.send_file("unsized.txt", size=None)
    assert STYLE_WARN_FILE_SIZE_UNKNOWN in env.last_edit
    await env.send_file("broken.docx", b"PK\x03\x04 broken")
    assert STYLE_WARN_FILE_UNREADABLE in env.last_edit

    assert env.downloaded == ["broken.docx"]  # nothing past the type and size checks
    assert "Получено: 0 знаков" in env.last_edit


# --- one Пространство's collection is never another's ---


async def test_bs_messages_and_done_never_touch_as_collection(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    question = env.question_id
    await env.send(PART * 3)

    assert not await env.send("чужой текст", env.b)
    assert await env.style.done(env.b, CHAT, USER, question) == STALE_BUTTON

    assert env.pending.rows[(CHAT, USER)].kind == STYLE_SAMPLES_KIND  # A's wait still open
    await env.done()
    (job,) = env.jobs.jobs.values()
    assert "чужой текст" not in job.payload["samples"]


# --- the result ---


async def _finished(env: Env, output: dict | None, ref: str = "a:s") -> int:
    await env.style.start(env.a, CHAT, USER, ref)
    await env.send(PART * 3)
    await env.done()
    (job_id,) = env.jobs.jobs
    result = env.jobs.finish(job_id, output)
    await env.style.forget_samples(result)
    await env.style.deliver(CHAT, result)
    return job_id


async def test_the_result_shows_the_profile_and_the_samples_are_gone(env: Env) -> None:
    job_id = await _finished(env, PROFILE_OUTPUT, "a:o")

    _, message_id, text, markup = env.bot.edited[-1]
    assert message_id == env.jobs.jobs[job_id].payload["message_id"]
    assert text.startswith("Вот как вы пишете:\n\nВы пишете коротко и по делу.")
    assert "Тон: прямой" in text
    assert button_data(markup) == [[f"yk:{job_id}"], [f"yf:{job_id}"], ["ya:a:o"]]
    assert set(env.jobs.jobs[job_id].payload) == {"source", "mode", "message_id"}
    assert env.store.values == {}  # nothing saved before «✅»


def test_failed_and_too_short_results_offer_other_samples_or_presets() -> None:
    text, markup = render_result(1, StyleRef("p", "o"), None)
    assert text == STYLE_FAILED and button_data(markup) == [["ya:p:o"], ["yp:o"]]

    text, _ = render_result(
        1, StyleRef("a"), {"status": "too_short", "chars": 900, "unread": ["https://t.me/x"]}
    )
    assert "Не хватило текста: 900 знаков" in text and "https://t.me/x" in text


async def test_confirm_saves_the_profile_and_offers_the_way_on(env: Env) -> None:
    job_id = await _finished(env, PROFILE_OUTPUT)
    env.onboarding.is_needed = True

    assert await env.style.confirm(env.a, CHAT, 500, str(job_id))

    profile = (await env.settings.read(WS_A)).style_profile
    assert profile is not None and profile.summary == "Вы пишете коротко и по делу."
    assert env.bot.edited[-1][1] == 500
    assert button_data(env.bot.edited[-1][3]) == [["ob:project"]]
    assert (await env.settings.read(WS_B)).style_profile is None


async def test_confirm_of_a_failed_or_foreign_job_saves_nothing(env: Env) -> None:
    job_id = await _finished(env, None)

    assert not await env.style.confirm(env.a, CHAT, 500, str(job_id))
    assert not await env.style.confirm(env.b, CHAT, 500, str(job_id))
    assert not await env.style.confirm(env.a, CHAT, 500, "nope")
    assert env.store.values == {}


async def test_fix_saves_corrections_and_a_later_ok_keeps_them(env: Env) -> None:
    job_id = await _finished(env, PROFILE_OUTPUT)

    assert await env.style.ask_fix(env.a, CHAT, USER, 500, str(job_id))
    assert env.bot.markups[-1] == (CHAT, 500, None)  # the result message loses its buttons
    assert env.pending.rows[(CHAT, USER)].kind == STYLE_FIX_KIND
    assert await env.style.handle_reply(env.a, CHAT, USER, "x" * 501)  # too long: asked again
    assert (await env.settings.read(WS_A)).style_profile is None

    assert await env.style.handle_reply(env.a, CHAT, USER, "Я пишу на «ты»")

    profile = (await env.settings.read(WS_A)).style_profile
    assert profile is not None and profile.corrections == "Я пишу на «ты»"
    assert "Ваши поправки: Я пишу на «ты»" in env.bot.sent[-1][1]
    assert not await env.style.handle_reply(env.a, CHAT, USER, "ещё")  # the wait is closed

    assert await env.style.confirm(env.a, CHAT, 500, str(job_id))  # a stale «✅» elsewhere
    kept = (await env.settings.read(WS_A)).style_profile
    assert kept is not None and kept.corrections == "Я пишу на «ты»"


def test_settings_show_the_owners_corrections() -> None:

    persona = next(setting for setting in SETTING_FIELDS if setting.key == "persona")
    profile = StyleProfile(traits=(("tone", "т"),), summary="Вы пишете.", corrections="на «ты»")
    current = OwnerSettings(niche="n", persona=None, custom_persona=None, style_profile=profile)

    assert persona.detail is not None
    assert persona.detail(current) == "по вашим текстам\nВы пишете.\nВаши поправки: на «ты»"


# --- the Лимит (ADR-0015) ---


async def test_the_analysis_is_reserved_as_one_free_job(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    await env.send(PART * 3)

    await env.done()

    assert env.limits.reserved == [(ANALYZE_STYLE_JOB, 1)]
    assert len(env.jobs.jobs) == 1


@pytest.mark.parametrize("refusal", [GenerationBusy(), DailyFreeJobLimitReached()])
async def test_a_refused_analysis_keeps_the_samples_and_the_wait(
    env: Env, refusal: Exception
) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    await env.send(PART * 3)
    env.limits.refusal = refusal  # type: ignore[assignment]

    alert = await env.done()

    assert alert == error_alert_text(refusal)
    assert env.jobs.jobs == {}
    assert env.pending.rows[(CHAT, USER)].kind == STYLE_SAMPLES_KIND
    assert STYLE_WORKING not in [text for _, _, text, _ in env.bot.edited]

    env.limits.refusal = None
    assert await env.done() is None
    assert next(iter(env.jobs.jobs.values())).payload["samples"] == [PART * 3]


async def test_a_refusal_after_typed_done_is_said_in_the_chat(env: Env) -> None:
    await env.style.start(env.a, CHAT, USER, "a:s")
    await env.send(PART * 3)
    env.limits.refusal = GenerationBusy()

    assert await env.send("готово")

    assert env.bot.sent[-1][1] == error_alert_text(GenerationBusy())
